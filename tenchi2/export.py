"""Export FC 天地を喰らうII 諸葛孔明伝 data to tenchi2/output/: CSV (UTF-8-BOM), XLSX and PNG."""
import argparse
import csv
import os
import sys
from collections import deque

import rom as R
from openpyxl import Workbook
from PIL import Image, ImageDraw

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.append(os.path.join(HERE, "..", "tenchi"))  # appended, so `rom` stays tenchi2/rom.py
import strips_page

SCREEN_H = R.SCREEN_ROWS * 16  # 240 px
KIND_NAMES = {0x80: "よろい", 0x81: "かぶと", 0x82: "たて", 0x88: "そうび"}


def write_csv(out: str, name: str, header: list[str], rows: list[list], book: Workbook) -> None:
    with open(os.path.join(out, name + ".csv"), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)
    ws = book.create_sheet(name)
    ws.append(header)
    for row in rows:
        ws.append(row)


def item_kind(r: R.Rom, iid: int) -> str:
    k = r.item_kind(iid)
    if 1 <= k <= 5:
        return r.weapon_class(k)
    return KIND_NAMES.get(k, "")


# ---------------------------------------------------------------- tables
def generals(r: R.Rom) -> tuple[list[str], list[list]]:
    header = ["id", "name", "kanji", "ちりょく", "ぶりょく", "すばやさ", "ぶき", "ちけい", "策略群", "章(+00 高4位)",
              "+00 低4位", "+01", "+05..+0A (hex)", "+0C", "+0D", "頭像列 (bank:row)"]
    rows = []
    for g in range(R.GENERAL_COUNT):
        b = r.general_record(g)
        rows.append([g, r.general_name(g), r.general_kanji(g), b[2], b[3], b[4], r.weapon_class(b[0x0C] & 7),
                     r.terrain(b[0x0B] >> 5), b[0x0B] & 0x0F, b[0] >> 4, b[0] & 0x0F, b[1], b[5:0x0B].hex(" "),
                     f"{b[0x0C]:02X}", b[0x0D], " ".join(f"{bk:02X}:{i:02X}" for bk, i in r.portrait_rows(g))])
    return header, rows


def items(r: R.Rom) -> tuple[list[str], list[list]]:
    header = ["id", "name", "price", "種類", "攻擊/防禦", "kind byte"]
    rows = []
    for i in range(8, R.ITEM_COUNT):
        price, name = r.item(i)
        k = r.item_kind(i)
        power = r.item_power(i) if k else ""
        rows.append([i, name, price, item_kind(r, i), power, f"{k:02X}"])
    return header, rows


def shops(r: R.Rom) -> tuple[list[str], list[list]]:
    header = ["map", "武器屋 list", "武器屋", "道具屋 list", "道具屋"]
    rows = []
    for m in range(R.MAP_COUNT):
        if r.map_info(m)[0] != 1:
            continue
        w, i = r.shops(m)
        names = [" / ".join(r.item(x)[1] for x in r.shop_list(k)) if k or m == 2 else "" for k in (w, i)]
        if any(names):
            rows.append([m, w, names[0], i, names[1]])
    return header, rows


def tactics(r: R.Rom) -> tuple[list[str], list[list]]:
    groups = {g: set(r.tactic_group(g)) for g in range(1, 11)}
    header = ["id", "name", "類別", "習得等級", "可用策略群"]
    rows = []
    for t in range(1, R.TACTIC_COUNT):
        lv = r.tactic_level(t)
        rows.append([t, r.tactic_name(t), R.TACTIC_KINDS.get(r.tactic_kind(t), r.tactic_kind(t)),
                     "" if lv == 0xFF else lv, " ".join(str(g) for g in groups if t in groups[g])])
    return header, rows


def links(r: R.Rom) -> tuple[list[str], list[list]]:
    owner = map_regions(r, r.links())
    header = ["kind", "offset", "from map (推定)", "from y screen", "from x screen", "from y cell", "from x cell", "to map",
              "to y screen", "to x screen", "to y cell", "to x cell", "flags", "script"]
    rows = []
    for k in r.links():
        fy, fys, fx, fxs = k["from"]
        ty, tys, tx, txs = k["to"]
        rows.append([k["kind"], f"{k['offset']:05X}", owner.get((fys, fxs), ""), fys, fxs, fy, fx, k["map"], tys, txs, ty, tx,
                     f"{k['flags']:02X}", "" if k["script"] is None else f"{k['script']:04X}"])
    return header, rows


# ---------------------------------------------------------------- pictures
def portraits(r: R.Rom, out: str) -> None:
    pal = [R.nes_rgb(c) for c in R.STATUS_PALETTE]
    d = os.path.join(out, "portraits")
    os.makedirs(d, exist_ok=True)
    cw, chh = 84, 102
    sheet = Image.new("RGB", (16 * cw, 16 * chh), (40, 40, 40))
    dr, f = ImageDraw.Draw(sheet), strips_page.font(10)
    for g in range(R.GENERAL_COUNT):
        im = Image.new("RGB", (48, 48))
        im.putdata([pal[c] for row in r.portrait(g) for c in row])
        im.save(os.path.join(d, f"{g:03d}.png"))
        x, y = (g % 16) * cw, (g // 16) * chh
        sheet.paste(im, (x + 18, y + 4))
        codes = r.general_kanji_codes(g)
        kx = x + 18 + (48 - 16 * len(codes)) // 2
        for n, (s, i) in enumerate(codes):
            for gy, row in enumerate(r.kanji_glyph(s, i)):
                for gx, v in enumerate(row):
                    if v:
                        sheet.putpixel((kx + n * 16 + gx, y + 54 + gy), (255, 255, 255))
        for n, t in enumerate((str(g), r.general_name(g))):
            dr.text((x + (cw - dr.textlength(t, font=f)) // 2, y + 72 + n * 13), t, fill=(200, 200, 200), font=f)
    sheet.save(os.path.join(out, "portraits_sheet.png"))


def strip_data(r: R.Rom) -> dict:
    """Data for strips_page: strips are the (bank, row index) pairs in use, mirrored like the ぶしょう screen.
    231 (junk record) and the all-blank faces (0, 109, 249-255) are left out; there is no palette choice."""
    gens = []
    for g in range(R.GENERAL_COUNT):
        if g != 231 and any(c for row in r.portrait(g) for c in row):
            kanji = [strips_page.glyph_hex([v for row in r.kanji_glyph(s, i) for v in row])
                     for s, i in r.general_kanji_codes(g)]
            zh = r.general_kanji(g).strip("\u3000")
            gens.append({"id": g, "name": zh or r.general_name(g), "zh": zh, "kana": r.general_name(g),
                         "rows": r.portrait_rows(g), "c": 0, "k": kanji})
    keys = sorted({v for g in gens for v in g["rows"]})
    index = {v: n for n, v in enumerate(keys)}
    strips = []
    for b, idx in keys:
        o = b * 0x2000 + idx * R.PORTRAIT_ROW_BYTES
        px = [[0] * 48 for _ in range(8)]
        for t in range(6):
            lo, hi = r.prg[o + t * 16:o + t * 16 + 8], r.prg[o + t * 16 + 8:o + t * 16 + 16]
            for y in range(8):
                for x in range(8):
                    px[y][t * 8 + x] = ((lo[y] >> (7 - x)) & 1) | (((hi[y] >> (7 - x)) & 1) << 1)
        strips.append("".join(str(c) for row in px for c in row[::-1]))
    for g in gens:
        g["p"] = [index[v] for v in g.pop("rows")]
    hexrgb = ["#" + "".join(f"{x:02x}" for x in R.nes_rgb(c)) for c in R.STATUS_PALETTE]
    return {"game": "tenchi2", "title": "天地を喰らうII", "group_names": ["兜", "官帽"],
            "fixed": hexrgb[:3], "colors": hexrgb[3:], "strips": strips,
            "labels": [f"{b:X}:{i:02X}" for b, i in keys], "generals": gens}


def kanji_sheets(r: R.Rom, out: str) -> None:
    for s, n in ((0, 256), (1, 64)):
        im = Image.new("L", (16 * 18, (n // 16) * 18), 255)
        for i in range(n):
            for y, row in enumerate(r.kanji_glyph(s, i)):
                for x, v in enumerate(row):
                    if v:
                        im.putpixel(((i % 16) * 18 + 1 + x, (i // 16) * 18 + 1 + y), 0)
        im.save(os.path.join(out, f"kanji{s + 1}.png"))


class MapPainter:
    def __init__(self, r: R.Rom):
        self.r, self.cache = r, {}

    def cell(self, m: int, t: int, cid: int) -> Image.Image:
        key = (m, t, cid)
        if key not in self.cache:
            r = self.r
            chr_set = r.map_info(m)[1]
            pal = [R.nes_rgb(c) for c in r.map_palette(m)]
            im = Image.new("RGB", (16, 16))
            im.putdata([pal[c] for row in r.cell_pixels(t, cid, r.bg_tiles(chr_set)) for c in row])
            self.cache[key] = im
        return self.cache[key]

    def screen(self, m: int, row: int, col: int) -> Image.Image:
        t = self.r.tileset(m, row)
        im = Image.new("RGB", (256, SCREEN_H))
        for y, line in enumerate(self.r.screen_cells(row, col, t)[:R.SCREEN_ROWS]):
            for x, cid in enumerate(line):
                im.paste(self.cell(m, t, cid), (x * 16, y * 16))
        return im


def uniform(r: R.Rom, row: int, col: int) -> bool:
    blocks = r.screen_blocks(*r.screen_entry(row, col))
    return len({b for line in blocks for b in line}) == 1


def map_regions(r: R.Rom, links: list[dict]) -> dict[tuple[int, int], int]:
    """Grid screen -> map number. Rows 68-90 are the world (map 0); the other screens go to the nearest link
    destination (multi-source BFS over non-blank screens of the same data bank)."""
    owner = {(row, col): 0 for row in R.WORLD_ROWS for col in range(R.GRID_W)}
    queue = deque()
    for k in links:
        s = (k["to"][1], k["to"][3])
        if k["map"] and s not in owner:
            owner[s] = k["map"]
            queue.append(s)
    while queue:
        y, x = queue.popleft()
        bank = r.screen_entry(y, x)[1]
        for s in ((y - 1, x), (y + 1, x), (y, x - 1), (y, x + 1)):
            if (0 <= s[0] < R.GRID_H and 0 <= s[1] < R.GRID_W and s not in owner
                    and r.screen_entry(*s)[1] == bank and not uniform(r, *s)):
                owner[s] = owner[(y, x)]
                queue.append(s)
    return owner


def world_view(r: R.Rom, links: list[dict]) -> set[tuple[int, int]]:
    """Grid cells (row * 15 + y, col * 16 + x) of the world the camera can show: walk (terrain < WALL) from every
    world cell a link starts at or leads to, then add the camera window around each reachable cell."""
    types = {}
    for row in R.WORLD_ROWS:
        t = r.tileset(0, row)
        for col in range(R.GRID_W):
            for y, line in enumerate(r.screen_cells(row, col, t)[:R.SCREEN_ROWS]):
                for x, c in enumerate(line):
                    types[(row * R.SCREEN_ROWS + y, col * 16 + x)] = r.cell_type(t, c)
    seen = set()
    for k in links:
        for (y, ys, x, xs), on_world in ((k["from"], True), (k["to"], k["map"] == 0)):
            c = (ys * R.SCREEN_ROWS + y, xs * 16 + x)
            if on_world and c in types:
                seen.add(c)
    queue = deque(seen)
    while queue:
        y, x = queue.popleft()
        for c in ((y - 1, x), (y + 1, x), (y, x - 1), (y, x + 1)):
            if c not in seen and types.get(c, R.WALL) < R.WALL:
                seen.add(c)
                queue.append(c)
    return {(y + dy, x + dx) for y, x in seen
            for dy in range(-R.VIEW_UP, R.VIEW_UP + 1) for dx in range(-R.VIEW_LEFT, R.VIEW_LEFT)}


def label(draw: ImageDraw.ImageDraw, x: int, y: int, text: str, colour: tuple[int, int, int]) -> None:
    draw.rectangle((x, y, x + 31, y + 15), outline=colour, width=2)
    draw.rectangle((x + 2, y - 11, x + 2 + 6 * len(text) + 3, y - 1), fill=(0, 0, 0))
    draw.text((x + 4, y - 11), text, fill=colour)


def maps(r: R.Rom, out: str) -> list[list]:
    """world.png, maps/map_NNN.png with link sources boxed and labelled with the destination map."""
    links = r.links()
    owner = map_regions(r, links)
    painter = MapPainter(r)
    by_map: dict[int, list[tuple[int, int]]] = {}
    for s, m in owner.items():
        by_map.setdefault(m, []).append(s)
    d = os.path.join(out, "maps")
    os.makedirs(d, exist_ok=True)
    rows = []
    view = world_view(r, links)
    by_map[0] = [s for s in by_map[0] if any((s[0] * R.SCREEN_ROWS + y, s[1] * 16 + x) in view
                                              for y in range(R.SCREEN_ROWS) for x in range(16))]
    for m, screens in sorted(by_map.items()):
        parts = components(screens) if m else [screens]
        for n, part in enumerate(parts):
            suffix = "" if len(parts) == 1 else "_" + chr(ord("a") + n)
            name = "world.png" if m == 0 else os.path.join("maps", f"map_{m:03d}{suffix}.png")
            rows.append(draw_map(r, painter, links, owner, m, part, os.path.join(out, name),
                                 view if m == 0 else None) + [name])
    return rows


def components(screens: list[tuple[int, int]]) -> list[list[tuple[int, int]]]:
    """4-connected groups of screens, largest first."""
    left, out = set(screens), []
    while left:
        todo, part = [left.pop()], []
        while todo:
            y, x = s = todo.pop()
            part.append(s)
            for t in ((y - 1, x), (y + 1, x), (y, x - 1), (y, x + 1)):
                if t in left:
                    left.remove(t)
                    todo.append(t)
        out.append(part)
    return sorted(out, key=lambda p: (-len(p), min(p)))


def draw_map(r: R.Rom, painter: MapPainter, links: list[dict], owner: dict, m: int,
             screens: list[tuple[int, int]], path: str, view: set[tuple[int, int]] | None = None) -> list:
    """view: grid cells the camera can show; the others are drawn at 30 % brightness."""
    y0, y1 = min(s[0] for s in screens), max(s[0] for s in screens)
    x0, x1 = min(s[1] for s in screens), max(s[1] for s in screens)
    im = Image.new("RGB", ((x1 - x0 + 1) * 256, (y1 - y0 + 1) * SCREEN_H))
    for y, x in screens:
        im.paste(painter.screen(m, y, x), ((x - x0) * 256, (y - y0) * SCREEN_H))
    if view is not None:
        mask = Image.new("L", im.size, 0)
        md = ImageDraw.Draw(mask)
        for cy, cx in view:
            px, py = (cx - x0 * 16) * 16, (cy - y0 * R.SCREEN_ROWS) * 16
            if y0 <= cy // R.SCREEN_ROWS <= y1 and x0 <= cx // 16 <= x1:
                md.rectangle((px, py, px + 15, py + 15), fill=255)
        im = Image.composite(im, Image.eval(im, lambda v: v * 3 // 10), mask)
    draw = ImageDraw.Draw(im)
    inside = set(screens)
    for k in links:
        fy, fys, fx, fxs = k["from"]
        if (fys, fxs) in inside and (view is None or (fys * R.SCREEN_ROWS + fy, fxs * 16 + fx) in view):
            colour = (255, 255, 0) if k["kind"] == "door" else (0, 255, 255)
            label(draw, (fxs - x0) * 256 + fx * 16, (fys - y0) * SCREEN_H + fy * 16, str(k["map"]), colour)
    im.save(path)
    area, chr_set = r.map_info(m)
    return [m, area, chr_set, f"{y0}-{y1}", f"{x0}-{x1}", len(screens)]


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--rom", default=R.DEFAULT_ROM)
    ap.add_argument("--out", default=os.path.join(HERE, "output"))
    a = ap.parse_args()
    os.makedirs(a.out, exist_ok=True)
    r = R.Rom(a.rom)
    text = R.Text(r)
    book = Workbook()
    book.remove(book.active)
    for name, (header, rows) in (("generals", generals(r)), ("items", items(r)), ("shops", shops(r)),
                                 ("tactics", tactics(r)), ("links", links(r))):
        write_csv(a.out, name, header, rows, book)
    write_csv(a.out, "exp", ["level", "exp"], [[lv + 2, e] for lv, e in enumerate(r.exp_table())], book)
    write_csv(a.out, "messages", ["id", "offset", "text"],
              [[m, f"{r.message_offset(m):05X}", text.message(m)] for m in range(R.MESSAGE_COUNT)], book)
    portraits(r, a.out)
    strip_rows = strips_page.write(strip_data(r), a.out)
    write_csv(a.out, "portrait_strips", list(strip_rows[0]), [list(x.values()) for x in strip_rows], book)
    kanji_sheets(r, a.out)
    write_csv(a.out, "maps", ["map", "area/tileset", "BG tile set", "grid rows", "grid cols", "screens", "png"],
              maps(r, a.out), book)
    book.save(os.path.join(a.out, "tenchi2.xlsx"))
    print("wrote", a.out)


if __name__ == "__main__":
    main()
